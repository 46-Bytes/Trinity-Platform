"""
Build the Sale Ready program fixtures from the client's program management Excel.

Writes, under files/sale_ready/:
  stages.json          the 15 stages (structure from the build brief)
  task_templates.json  must-do and optional tasks per stage
  dd_templates.json    the 210-item due diligence checklist, mapped to stages
  sale_planner.json    Sale Planner option lists, questions and issues
  templates.json       the Templates library (names only)
  unresolved.json      content held back until the client answers (C-numbers)

Keys already present in an existing fixture are preserved when titles match,
so regenerating after a wording change does not re-key rows engagements use.

Usage (from backend/):
    python scripts/build_sale_ready_fixtures.py --xlsx "<path to the program sheet>.xlsx"
    python scripts/build_sale_ready_fixtures.py --xlsx <xlsx> --mockup <mockup.html>   # cross-check
    python scripts/build_sale_ready_fixtures.py --xlsx <xlsx> --check                  # verify only
"""
import argparse
import hashlib
import json
import os
import re
import sys
from collections import Counter, OrderedDict

import openpyxl

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(BACKEND_DIR, "files", "sale_ready")
PROGRAM_TYPE = "sale_ready"

# The 15 stages and their order come from the build brief ("The structure").
# Titles follow the brief; display codes P1-P7 follow the mockup.
STAGES = [
    # code, stage_type, display_code, title, source sheet, task_creation, ui_variant
    ("DIAG", "pre_module", "P1", "Diagnostic", "Diagnostic", "on_engagement_create", None),
    ("APPRAISAL", "pre_module", "P2", "Appraisal", "Appraisal", "on_engagement_create", None),
    ("PRIORITISE", "pre_module", "P3", "Prioritisation & Roadmap", "Prioritisation", "on_engagement_create", None),
    ("WORKSHOP", "pre_module", "P4", "Planning Workshop", "Planning Workshop", "on_engagement_create", None),
    ("M1", "module", "M1", "Financial Clarity & Reporting", "Module 1 - Financials", "on_start", None),
    # M2 takes the Diagnostic's name (C3, resolved 21 Sep 2026); the Excel and the
    # brief both call it "Legal, Compliance, Property & Assets".
    ("M2", "module", "M2", "Legal, Compliance & Property", "Module 2 - Legal, Property, Ass", "on_start", None),
    ("M3", "module", "M3", "Owner Dependency & Operations", "Module 3 -Owner & Operations", "on_start", None),
    ("M4", "module", "M4", "People", "Module 4 - People", "on_start", None),
    ("M5", "module", "M5", "Customer, Product & Revenue Quality", "Module 5 - Customer & Product", "on_start", None),
    ("M6", "module", "M6", "Brand, IP & Intangibles", "Module 6 - Brand & IP", "on_start", None),
    ("M7", "module", "M7", "Tax, Compliance & Regulatory", "Module 7 - Tax & Compliance", "on_start", None),
    ("M8", "module", "M8", "Due Diligence Preparation", "Module 8 - DD Preparation", "on_start", None),
    ("TRANSITION", "post_module", "P5", "Transition & Handover Planning", "Transition & Handover Planning", "on_engagement_create", None),
    ("SALE_PLANNER", "post_module", "P6", "Sale Planner Checklist", "Sale Planner Checklist", "on_engagement_create", "sale_planner"),
    ("CLOSEOUT", "post_module", "P7", "Re-appraisal / Listing / Close-out", None, "on_engagement_create", "closeout"),
]

DD_SHEET = "Due Diligence Checklist"
EXPECTED_DD_ITEMS = 210
EXPECTED_DD_CATEGORIES = 16

# The mockup's stage keys, for the optional cross-check only.
MOCKUP_STAGE_KEYS = {"appraisal": "APPRAISAL", "transition": "TRANSITION",
                     **{f"M{i}": f"M{i}" for i in range(1, 9)}}

# Rows 2-25 of the Templates sheet are the library; rows 26+ look like working notes (C18).
TEMPLATE_ROWS = range(2, 26)


class FixtureBuildError(Exception):
    pass


def clean(value):
    """Cell text, trimmed per line, or None for an empty cell."""
    if value is None:
        return None
    text = "\n".join(line.strip() for line in str(value).strip().splitlines())
    text = re.sub(r"[ \t]+", " ", text).strip()
    return text or None


def norm(text):
    return re.sub(r"\s+", " ", (text or "")).strip().lower()


def split_code(label):
    """'3.1 Historical Financial Statements' -> ('3.1', 'Historical Financial Statements')."""
    match = re.match(r"^(\d+(?:\.\d+)?)\.?\s+(.*)$", label or "")
    if not match:
        raise FixtureBuildError(f"No leading code in {label!r}")
    return match.group(1), match.group(2).strip()


def load_existing(name):
    path = os.path.join(OUT_DIR, name)
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)


# ----------------------------------------------------------------------
# Tasks
# ----------------------------------------------------------------------
def classify_section(header):
    """Map a 🔹/🔸 section header to (section, group_title)."""
    text = header.lstrip("🔹🔸 ").strip()
    lowered = text.lower()
    if lowered.startswith("must-do"):
        return "must_do", None
    if lowered.startswith("optional"):
        return "optional", None
    if "client specific" in lowered:
        return "client_specific", None
    section_match = re.match(r"^section\s+\d+:\s*(.+)$", text, re.IGNORECASE)
    if section_match:
        # Phase sheets group their tasks under numbered sections; all of them are must-do.
        return "must_do", section_match.group(1).strip()
    raise FixtureBuildError(f"Unrecognised section header {header!r}")


def parse_tasks(ws):
    """Tasks from column A of a phase or module sheet, stopping at the register block."""
    tasks, ignored = [], []
    section = group_title = None
    for row in range(1, ws.max_row + 1):
        text = clean(ws.cell(row=row, column=1).value)
        if not text:
            continue
        if text == "Document Name":
            break
        if text.startswith(("🔹", "🔸")):
            section, group_title = classify_section(text)
            continue
        if section is None:
            continue  # sheet header rows above the first section
        if section == "client_specific":
            # Client-specific tasks are added per engagement, never templated.
            ignored.append({"row": row, "text": text})
            continue
        tasks.append({"section": section, "group_title": group_title, "title": text, "row": row})
    return tasks, ignored


def build_task_templates(wb, existing):
    previous = {}
    for item in (existing or {}).get("items", []):
        previous[(item["stage_code"], item["section"], norm(item["title"]))] = item["template_key"]

    items, ignored_by_stage = [], {}
    for code, _type, _disp, _title, sheet, _creation, _variant in STAGES:
        if not sheet or sheet not in wb.sheetnames or code in ("TRANSITION", "SALE_PLANNER"):
            continue
        tasks, ignored = parse_tasks(wb[sheet])
        if ignored:
            ignored_by_stage[code] = ignored
        used = {k for (s, _sec, _t), k in previous.items() if s == code}
        counters = Counter()
        for order, task in enumerate(tasks, start=1):
            short = "MUST" if task["section"] == "must_do" else "OPT"
            key = previous.get((code, task["section"], norm(task["title"])))
            if key is None:
                while True:
                    counters[short] += 1
                    key = f"{code}-{short}-{counters[short]:02d}"
                    if key not in used:
                        break
            used.add(key)
            items.append(OrderedDict([
                ("program_type", PROGRAM_TYPE),
                ("stage_code", code),
                ("template_key", key),
                ("section", task["section"]),
                ("group_title", task["group_title"]),
                ("title", task["title"]),
                ("default_order", order),
                ("source_row", task["row"]),
            ]))

    items.extend(closeout_tasks())
    return items, ignored_by_stage


# The Excel has no close-out sheet; these four come from the mockup, which the
# brief makes the spec ("the last phase has the close-out tasks"). Emitted here
# so regenerating from the Excel does not drop them (C2).
CLOSEOUT_TASKS = [
    "Review the final position against the original appraisal",
    "Decide whether a fresh appraisal is required and, if so, run it",
    "Refer the client to Benchmark Business Sales for listing",
    "Agree the ongoing assistance arrangement",
]


def closeout_tasks():
    return [
        OrderedDict([
            ("program_type", PROGRAM_TYPE),
            ("stage_code", "CLOSEOUT"),
            ("template_key", f"CLOSEOUT-MUST-{i:02d}"),
            ("section", "must_do"),
            ("group_title", "Close-out"),
            ("title", title),
            ("default_order", i),
            ("source_row", None),
        ])
        for i, title in enumerate(CLOSEOUT_TASKS, start=1)
    ]


# ----------------------------------------------------------------------
# Due diligence
# ----------------------------------------------------------------------
def parse_master_dd(ws):
    """The 210 items from the master sheet. Rows with only an action step continue the item above."""
    header_row = next(
        r for r in range(1, 10) if clean(ws.cell(row=r, column=1).value) == "Category"
    )
    items, category, sub_item = [], None, None
    for row in range(header_row + 1, ws.max_row + 1):
        cat = clean(ws.cell(row=row, column=1).value)
        sub = clean(ws.cell(row=row, column=2).value)
        doc = clean(ws.cell(row=row, column=3).value)
        action = clean(ws.cell(row=row, column=4).value)
        category = cat or category
        sub_item = sub or sub_item
        if doc:
            items.append({"category": category, "sub_item": sub_item, "document": doc,
                          "action_steps": [action] if action else [], "row": row})
        elif action and items:
            items[-1]["action_steps"].append(action)
    return items


def parse_sheet_dd(ws):
    """(sub_item_code, document) pairs listed on a phase or module sheet."""
    for row in range(1, 15):
        for col in range(1, ws.max_column + 1):
            if clean(ws.cell(row=row, column=col).value) == "Document/Information Required":
                doc_col, sub_col, cat_col = col, col - 1, col - 2
                pairs, sub_item = [], None
                for r in range(row + 1, ws.max_row + 1):
                    # Where the DD block starts in column A, the register header ends it.
                    if clean(ws.cell(row=r, column=cat_col).value) == "Document Name":
                        break
                    sub = clean(ws.cell(row=r, column=sub_col).value)
                    doc = clean(ws.cell(row=r, column=doc_col).value)
                    if sub and re.match(r"^\d+\.\d+\s", sub):
                        sub_item = sub
                    if doc and sub_item:
                        pairs.append((split_code(sub_item)[0], norm(doc)))
                return pairs
    return []


def build_dd_templates(wb, existing):
    master = parse_master_dd(wb[DD_SHEET])

    stage_of = {}
    for code, _type, _disp, _title, sheet, _creation, _variant in STAGES:
        if not sheet or sheet not in wb.sheetnames:
            continue
        for pair in parse_sheet_dd(wb[sheet]):
            if pair in stage_of and stage_of[pair] != code:
                raise FixtureBuildError(f"DD item {pair} listed on {stage_of[pair]} and {code}")
            stage_of[pair] = code

    previous = {(i["sub_item_code"], norm(i["document"])): i["item_key"]
                for i in (existing or {}).get("items", [])}
    used = set(previous.values())
    next_number = 0

    items, unmapped, wording_diffs = [], [], []
    for order, raw in enumerate(master, start=1):
        cat_code, cat_label = split_code(raw["category"])
        sub_code, sub_label = split_code(raw["sub_item"])
        pair = (sub_code, norm(raw["document"]))
        stage = stage_of.get(pair)
        if stage is None:
            # Master wording can extend a module sheet's (or the reverse); accept one unique prefix match.
            candidates = [(p, s) for p, s in stage_of.items()
                          if p[0] == sub_code and (pair[1].startswith(p[1]) or p[1].startswith(pair[1]))]
            if len(candidates) == 1:
                stage = candidates[0][1]
                wording_diffs.append({"sub_item_code": sub_code, "master": raw["document"],
                                      "stage_sheet": candidates[0][0][1], "stage_code": stage})
        if stage is None:
            unmapped.append({"row": raw["row"], "sub_item": raw["sub_item"], "document": raw["document"]})
            continue
        key = previous.get(pair)
        if key is None:
            while True:
                next_number += 1
                key = f"DD-{next_number:03d}"
                if key not in used:
                    break
            used.add(key)
        items.append(OrderedDict([
            ("program_type", PROGRAM_TYPE),
            ("item_key", key),
            ("stage_code", stage),
            ("category_code", cat_code),
            ("category", cat_label),
            ("sub_item_code", sub_code),
            ("sub_item", sub_label),
            ("document", raw["document"]),
            ("action_step", "\n".join(raw["action_steps"]) or None),
            ("default_order", order),
            ("source_row", raw["row"]),
        ]))

    if unmapped:
        raise FixtureBuildError(f"{len(unmapped)} DD items are on no stage sheet: {unmapped[:5]}")
    return items, wording_diffs


# ----------------------------------------------------------------------
# Sale Planner, templates
# ----------------------------------------------------------------------
SALE_PLANNER_HEADINGS = {
    "Type of Sale": "sale_types",
    "Structure of Sale": "sale_structures",
    "Best Value Proposition": "value_proposition_structures",
    "Marketing Plan": "marketing_questions",
}


def slug(text):
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def build_sale_planner(ws):
    issues_col = next(
        c for c in range(1, ws.max_column + 1) if clean(ws.cell(row=1, column=c).value) == "Issues To Be Addressed"
    )
    lists = {key: [] for key in SALE_PLANNER_HEADINGS.values()}
    current = None
    for row in range(1, ws.max_row + 1):
        text = clean(ws.cell(row=row, column=1).value)
        if not text:
            continue
        if text in SALE_PLANNER_HEADINGS:
            current = SALE_PLANNER_HEADINGS[text]
            continue
        if current:
            lists[current].append(text)
    issues = [clean(ws.cell(row=r, column=issues_col).value) for r in range(2, ws.max_row + 1)]
    issues = [i for i in issues if i]

    def keyed(values):
        return [{"key": slug(v), "label": v} for v in values]

    return OrderedDict([
        ("sale_types", keyed(lists["sale_types"])),
        ("sale_structures", keyed(lists["sale_structures"])),
        ("value_proposition_structures", keyed(lists["value_proposition_structures"])),
        ("marketing_questions", keyed(lists["marketing_questions"])),
        ("issues", keyed(issues)),
    ])


def build_templates(ws):
    library, notes = [], []
    for row in range(2, ws.max_row + 1):
        text = next((clean(c.value) for c in ws[row] if clean(c.value)), None)
        if not text:
            continue
        (library if row in TEMPLATE_ROWS else notes).append((row, text))
    items = [OrderedDict([("program_type", PROGRAM_TYPE), ("template_key", slug(text)),
                          ("name", text), ("display_order", i)])
             for i, (_row, text) in enumerate(library, start=1)]
    return items, [{"row": r, "text": t} for r, t in notes]


# ----------------------------------------------------------------------
# Assembly
# ----------------------------------------------------------------------
def build(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    with open(xlsx_path, "rb") as f:
        source = OrderedDict([("file", os.path.basename(xlsx_path)),
                              ("sha256", hashlib.sha256(f.read()).hexdigest())])

    stages = [OrderedDict([
        ("program_type", PROGRAM_TYPE), ("stage_code", code), ("stage_type", stage_type),
        ("display_code", display), ("title", title), ("default_order", order),
        ("task_creation", creation), ("ui_variant", variant), ("source_sheet", sheet),
    ]) for order, (code, stage_type, display, title, sheet, creation, variant) in enumerate(STAGES, start=1)]

    tasks, ignored_client_rows = build_task_templates(wb, load_existing("task_templates.json"))
    dd, dd_wording_diffs = build_dd_templates(wb, load_existing("dd_templates.json"))
    planner = build_sale_planner(wb["Sale Planner Checklist"])
    library, template_notes = build_templates(wb["Templates"])

    onboarding, _ = parse_tasks(wb["Onboarding"])
    named_people = [t["template_key"] for t in tasks if re.search(r"\b(Pete|Marc)\b", t["title"])]
    sheet5 = [clean(c.value) for r in wb["Sheet5"].iter_rows() for c in r if clean(c.value)]

    unresolved = OrderedDict([
        ("C1_onboarding_tasks", {
            "note": "Onboarding is not one of the brief's 15 stages; these tasks are not seeded.",
            "tasks": [{"section": t["section"], "group_title": t["group_title"], "title": t["title"]}
                      for t in onboarding]}),
        ("C2_closeout_tasks", {
            "note": "RESOLVED 21 Sep 2026: the Excel has no close-out tasks; the brief makes the mockup the "
                    "spec, so the mockup's four are emitted by closeout_tasks() and seeded.",
            "tasks": CLOSEOUT_TASKS}),
        ("C2_named_people_in_tasks", {
            "note": "Seeded verbatim from the Excel; the mockup rewords them without names.",
            "template_keys": named_people}),
        ("C2_closeout_title", {
            "note": "Brief: 'Re-appraisal / Listing / Close-out'. Mockup: 'Re-appraisal, Listing & Close-out'.",
            "seeded": "Re-appraisal / Listing / Close-out"}),
        ("C2_dd_wording_differences", {
            "note": "Master checklist wording (seeded) differs from the stage sheet's wording.",
            "items": dd_wording_diffs}),
        ("C3_m2_title", {
            "note": "RESOLVED 21 Sep 2026: use the Diagnostic's name. Note the Excel and the brief both say "
                    "'Legal, Compliance, Property & Assets'.",
            "seeded_stage_title": "Legal, Compliance & Property",
            "excel_and_brief_title": "Legal, Compliance, Property & Assets"}),
        ("C4_stage_guide_text", {
            "note": "Purpose / how it runs / watch for / per-stage templates exist only in the mockup; not extracted."}),
        ("C9_register_renewal_cost", {
            "note": "Excel register has Renewal Cost; the mockup register does not. Register not seeded."}),
        ("C18_templates_sheet_notes", {
            "note": "Templates sheet rows after the 24-item library; not seeded.", "rows": template_notes}),
        ("C19_sheet5", {"note": "Sheet5 content; not seeded.", "cells": sheet5}),
        ("ignored_client_specific_rows", {
            "note": "Text found under 'Additional Client Specific Tasks'; client-specific tasks are never templated.",
            "rows": ignored_client_rows}),
    ])

    def wrap(items):
        return OrderedDict([("source", source), ("items", items)])

    return OrderedDict([
        ("stages.json", wrap(stages)),
        ("task_templates.json", wrap(tasks)),
        ("dd_templates.json", wrap(dd)),
        ("sale_planner.json", OrderedDict([("source", source), ("config", planner)])),
        ("templates.json", wrap(library)),
        ("unresolved.json", OrderedDict([("source", source), ("items", unresolved)])),
    ])


def validate(outputs):
    dd = outputs["dd_templates.json"]["items"]
    problems = []
    if len(dd) != EXPECTED_DD_ITEMS:
        problems.append(f"expected {EXPECTED_DD_ITEMS} DD items, got {len(dd)}")
    if len({i["category_code"] for i in dd}) != EXPECTED_DD_CATEGORIES:
        problems.append("expected 16 DD categories")
    if len({i["item_key"] for i in dd}) != len(dd):
        problems.append("duplicate DD item_key")
    tasks = outputs["task_templates.json"]["items"]
    keys = [(t["stage_code"], t["template_key"]) for t in tasks]
    if len(set(keys)) != len(keys):
        problems.append("duplicate task template_key within a stage")
    if len(outputs["templates.json"]["items"]) != 24:
        problems.append("expected 24 library templates")
    if problems:
        raise FixtureBuildError("; ".join(problems))


def cross_check_mockup(outputs, mockup_path):
    """Compare stage mapping and counts with the mockup's embedded DATA."""
    with open(mockup_path, encoding="utf-8") as f:
        html = f.read()
    match = re.search(r"const DATA=(\{.*?\});\s*\nconst el=", html, re.S)
    if not match:
        raise FixtureBuildError("Could not find DATA in the mockup")
    data = json.loads(match.group(1))
    ours = {(i["sub_item_code"], norm(i["document"])): i["stage_code"] for i in outputs["dd_templates.json"]["items"]}
    mismatches = []
    for item in data["dd"]:
        pair = (split_code(item["sub"])[0], norm(item["doc"]))
        expected = MOCKUP_STAGE_KEYS.get(item["mod"])
        if ours.get(pair) != expected:
            mismatches.append((pair, ours.get(pair), expected))
    if mismatches:
        raise FixtureBuildError(f"{len(mismatches)} DD stage mappings differ from the mockup: {mismatches[:5]}")
    return len(data["dd"])


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", required=True, help="Path to the Sale Ready program management Excel")
    parser.add_argument("--mockup", help="Optional: mockup HTML to cross-check DD stage mapping")
    parser.add_argument("--check", action="store_true", help="Build and validate, write nothing")
    args = parser.parse_args()

    try:
        outputs = build(args.xlsx)
        validate(outputs)
        if args.mockup:
            print(f"Mockup cross-check: {cross_check_mockup(outputs, args.mockup)} DD items agree")
    except FixtureBuildError as exc:
        print(f"FAILED: {exc}", file=sys.stderr)
        sys.exit(1)

    dd = outputs["dd_templates.json"]["items"]
    tasks = outputs["task_templates.json"]["items"]
    print(f"stages={len(outputs['stages.json']['items'])} task_templates={len(tasks)} "
          f"(must_do={sum(t['section'] == 'must_do' for t in tasks)}, "
          f"optional={sum(t['section'] == 'optional' for t in tasks)}) dd_items={len(dd)} "
          f"templates={len(outputs['templates.json']['items'])}")
    print("DD per stage:", dict(Counter(i["stage_code"] for i in dd)))

    if args.check:
        print("CHECK ONLY - nothing written.")
        return
    os.makedirs(OUT_DIR, exist_ok=True)
    for name, payload in outputs.items():
        with open(os.path.join(OUT_DIR, name), "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2, ensure_ascii=False)
            f.write("\n")
    print(f"Wrote {len(outputs)} files to {OUT_DIR}")


if __name__ == "__main__":
    main()
